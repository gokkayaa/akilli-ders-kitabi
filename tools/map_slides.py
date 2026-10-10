import sys, json, re, io, numpy as np, cv2
from pptx import Presentation
from pptx.util import Emu
W=sys.argv[2]; p=Presentation(sys.argv[1]); SW,SH=p.slide_width,p.slide_height
CLS={'7A','7B','7C','7D','7E','7F','8A','8B','8C'}
sift=cv2.SIFT_create()
pagecache={}
def page(key):
    if key not in pagecache:
        b,n=key
        im=cv2.imread(f'{W}/png/{b}-{n:03d}-c.png',cv2.IMREAD_GRAYSCALE)
        k,d=sift.detectAndCompute(im,None); pagecache[key]=(im.shape,k,d)
    return pagecache[key]
bf=cv2.BFMatcher()
def allshapes(shs):
    for sh in shs:
        if sh.shape_type==6: yield from allshapes(sh.shapes)
        else: yield sh
def rect(sh): return [sh.left/SW*100, sh.top/SH*100, sh.width/SW*100, sh.height/SH*100]
res=[]
for si,s in enumerate(p.slides,1):
    pass
    shs=list(allshapes(s.shapes))
    labels=[]
    for sh in shs:
        if sh.has_text_frame:
            t=sh.text_frame.text.strip()
            if re.fullmatch(r'\d+(-\d+)?',t) and (sh.left/SW>0.8) and sh.top/SH<0.15: labels+= [int(x) for x in t.split('-')]
    if not labels: continue
    cands=[(b,n) for n in set(labels) for b in ('dk2','dbt') if (b=='dk2' and 7<=n<=42) or (b=='dbt' and 93<=n<=108)]
    # big pictures = screenshots
    shots=[]
    for sh in shs:
        if sh.shape_type!=13: continue
        r=rect(sh)
        if r[2]*r[3] < 600: continue
        try: blob=sh.image.blob
        except Exception: continue
        img=cv2.imdecode(np.frombuffer(blob,np.uint8),cv2.IMREAD_GRAYSCALE)
        if img is None: continue
        k1,d1=sift.detectAndCompute(img,None)
        if d1 is None: continue
        best=None
        for c in cands:
            shp,k2,d2=page(c)
            m=bf.knnMatch(d1,d2,k=2); good=[a for a,*bb in m if bb and a.distance<0.7*bb[0].distance]
            if len(good)<12: continue
            src=np.float32([k1[g.queryIdx].pt for g in good]); dst=np.float32([k2[g.trainIdx].pt for g in good])
            A,inl=cv2.estimateAffinePartial2D(src,dst,ransacReprojThreshold=5)
            if A is None: continue
            ni=int(inl.sum())
            if not best or ni>best[0]: best=(ni,c,A,shp,img.shape,sh)
        if best and best[0]>=25:
            ni,c,A,shp,ishape,sh=best
            cl,ct,cr,cb=sh.crop_left,sh.crop_top,sh.crop_right,sh.crop_bottom
            shots.append({'shape':sh.name,'page':c,'A':A,'pshape':shp,'ishape':ishape,'r':rect(sh),'crop':(cl,ct,cr,cb),'inl':ni,'shobj':sh})
    if not shots: continue
    def to_page(sx,sy):
        # find containing shot (topmost last)
        for sh in reversed(shots):
            x,y,w,h=sh['r']
            if x-1<=sx<=x+w+1 and y-1<=sy<=y+h+1:
                ih,iw=sh['ishape']; cl,ct,cr,cb=sh['crop']
                u=(cl+(sx-x)/w*(1-cl-cr))*iw; v=(ct+(sy-y)/h*(1-ct-cb))*ih
                X,Y=sh['A']@np.array([u,v,1.0]); PH,PW=sh['pshape']
                return sh['page'], X/PW*100, Y/PH*100
        return None
    shotnames={sh['shape'] for sh in shots}
    for sh in shs:
        if sh.name in shotnames: continue
        r=rect(sh); item=None
        if sh.has_text_frame:
            t=sh.text_frame.text.strip().replace('\x0b',' ')
            if not t or t in CLS or 'bitti' in t or 'BİTTİ' in t or re.fullmatch(r'\d+(-\d+)?',t) and sh.left/SW>0.8: continue
            item={'kind':'text','text':t}
        elif sh.shape_type==13:
            # small answer image -> crop alpha bbox and save
            try:
                from PIL import Image
                im=Image.open(io.BytesIO(sh.image.blob)).convert('RGBA'); a=np.array(im)[:,:,3]
                ys,xs=np.where(a>20)
                if len(xs)==0: continue
                ih,iw=a.shape
                if r[2]*r[3] < 30 and r[0]<5 and r[1]>90: continue  # izo button
                r=[r[0]+xs.min()/iw*r[2], r[1]+ys.min()/ih*r[3], (xs.max()-xs.min())/iw*r[2], (ys.max()-ys.min())/ih*r[3]]
                fn=f'S{si}_{sh.shape_id}.png'; im.crop((xs.min(),ys.min(),xs.max()+1,ys.max()+1)).save(f'{W}/ansimg/{fn}')
                item={'kind':'img','file':fn}
            except Exception as e: continue
        elif sh.shape_type in (9,):  # line / connector
            x1,y1=sh.begin_x/SW*100, sh.begin_y/SH*100; x2,y2=sh.end_x/SW*100, sh.end_y/SH*100
            a=to_page(x1,y1); b=to_page(x2,y2)
            if a and b: res.append({'slide':si,'page':list(a[0]),'kind':'line','from':[round(a[1],2),round(a[2],2)],'to':[round(b[1],2),round(b[2],2)]})
            continue
        elif sh.shape_type==1:
            try: geom=sh.auto_shape_type
            except Exception: geom=None
            item={'kind':'shape','geom':str(geom)}
        else: continue
        a=to_page(r[0],r[1]); b=to_page(r[0]+r[2],r[1]+r[3])
        if not a or not b: continue
        item.update({'slide':si,'page':list(a[0]),'rect':[round(a[1],2),round(a[2],2),round(b[1]-a[1],2),round(b[2]-a[2],2)]})
        res.append(item)
    # baked red answers: red pixels in screenshot
    for sh in shots:
        img=cv2.imdecode(np.frombuffer(sh['shobj'].image.blob,np.uint8),cv2.IMREAD_COLOR)
        bch,g,rch=[img[:,:,i].astype(int) for i in range(3)]
        red=((rch>170)&(g<90)&(bch<90)).astype(np.uint8)
        # compare with page: page red at mapped location
        PH,PW=sh['pshape']; pc=cv2.imread(f"{W}/png/{sh['page'][0]}-{sh['page'][1]:03d}-c.png")
        warp=cv2.warpAffine(pc,cv2.invertAffineTransform(sh['A']),(img.shape[1],img.shape[0]),borderValue=(255,255,255))
        pb,pg_,pr=[warp[:,:,i].astype(int) for i in range(3)]
        pred=((pr>170)&(pg_<90)&(pb<90)).astype(np.uint8)
        pred=cv2.dilate(pred,np.ones((7,7),np.uint8))
        diff=red&(1-pred)
        diff=cv2.dilate(diff,np.ones((9,25),np.uint8))
        k,lab,st,cen=cv2.connectedComponentsWithStats(diff,8)
        for i in range(1,k):
            x,y,w,h,area=st[i]
            if area<300: continue
            pts=np.array([[x,y,1],[x+w,y+h,1]],float)@sh['A'].T
            fn=f"S{si}_red{i}_{sh['shape'].replace(' ','')}.png"
            cv2.imwrite(f'{W}/ansimg/{fn}',img[y:y+h,x:x+w])
            res.append({'slide':si,'page':list(sh['page']),'kind':'red','file':fn,'rect':[round(pts[0][0]/PW*100,2),round(pts[0][1]/PH*100,2),round((pts[1][0]-pts[0][0])/PW*100,2),round((pts[1][1]-pts[0][1])/PH*100,2)]})
    print(si, [ (s_['page'],s_['inl']) for s_ in shots], flush=True)
json.dump(res,open(f'{W}/pptx_answers.json','w'),ensure_ascii=False,indent=0)
print('items',len(res))
